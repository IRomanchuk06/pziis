#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Сборка отчётов по лабораторным работам (docx, стандарт БГУИР) из report/report.md.

Использование:
    python3 build_reports.py [lab_N ...]   # без аргументов — все найденные lab_N

Титульные данные берутся из student_info.json (заглушки вида [Фамилия И. О.]
подсвечиваются красным, чтобы их было видно в готовом документе).
"""
import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE = Path(__file__).resolve().parent.parent  # .../pziis
BUILD = Path(__file__).resolve().parent

LAB_TITLES = {
    1: "Оценка защищённости разработанного приложения",
    2: "Оценка защищённости сторонних программных систем",
    3: "Анализ безопасности кода",
    4: "Генерация случайных чисел",
    5: "Управление ключами",
    6: "Применение OpenSSL",
    7: "Методика оценки безопасности кода",
    8: "Управление доступом к приложениям и системам баз данных",
    9: "Управление доступом к объектам операционной системы Windows",
    10: "Управление доступом к объектам операционной системы Linux",
    11: "Анализ запросов из сети Интернет",
    12: "Совершенствование безопасности удалённых серверов",
}

PLACEHOLDER_RE = re.compile(r"^\[[^\]]+\]$")

BODY_FONT = "Times New Roman"
CODE_FONT = "Courier New"
BLACK = RGBColor(0, 0, 0)
PLACEHOLDER_COLOR = RGBColor(0xC0, 0x00, 0x00)
GRAY = RGBColor(0x60, 0x60, 0x60)


def is_placeholder(text: str) -> bool:
    return bool(PLACEHOLDER_RE.match(text.strip()))


def add_run(par, text, *, bold=False, italic=False, size=14, font=BODY_FONT,
            color=BLACK, mono=False):
    run = par.add_run(text)
    run.bold = bold
    run.italic = italic
    f = run.font
    f.name = CODE_FONT if mono else font
    f.size = Pt(size)
    f.color.rgb = color
    # попадание в themeFonts для кириллицы
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(attr), f.name)
    return run


TOKEN_RE = re.compile(r"(\*\*(?P<bold>.+?)\*\*|`(?P<code>[^`]+)`|\*(?P<ital>[^*\n]+)\*)", re.S)


def add_inline(par, text, *, size=14, bold=False, italic=False):
    pos = 0
    for m in TOKEN_RE.finditer(text):
        if m.start() > pos:
            add_run(par, text[pos:m.start()], bold=bold, italic=italic, size=size)
        if m.group("bold") is not None:
            add_inline(par, m.group("bold"), size=size, bold=True, italic=italic)
        elif m.group("code") is not None:
            add_run(par, m.group("code"), mono=True, size=max(size - 3, 9))
        else:
            add_inline(par, m.group("ital"), size=size, bold=bold, italic=True)
        pos = m.end()
    if pos < len(text):
        add_run(par, text[pos:], bold=bold, italic=italic, size=size)


def setup_styles(doc):
    st = doc.styles["Normal"]
    st.font.name = BODY_FONT
    st.font.size = Pt(14)
    st.font.color.rgb = BLACK
    rpr = st.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), BODY_FONT)
    pf = st.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(0)
    pf.space_before = Pt(0)

    for name, sz, before, after in (("Heading 1", 16, 18, 10),
                                    ("Heading 2", 14, 14, 8),
                                    ("Heading 3", 14, 12, 6)):
        h = doc.styles[name]
        h.font.name = BODY_FONT
        h.font.size = Pt(sz)
        h.font.bold = True
        h.font.italic = False
        h.font.color.rgb = BLACK
        hrpr = h.element.get_or_add_rPr()
        hrfonts = hrpr.find(qn("w:rFonts"))
        if hrfonts is None:
            hrfonts = OxmlElement("w:rFonts")
            hrpr.append(hrfonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            hrfonts.set(qn(attr), BODY_FONT)
        hpf = h.paragraph_format
        hpf.space_before = Pt(before)
        hpf.space_after = Pt(after)
        hpf.line_spacing = 1.15
        hpf.keep_with_next = True
        hpf.first_line_indent = Cm(0)


def setup_page(doc):
    sec = doc.sections[0]
    sec.page_height = Cm(29.7)
    sec.page_width = Cm(21.0)
    sec.top_margin = Cm(2)
    sec.bottom_margin = Cm(2)
    sec.left_margin = Cm(3)
    sec.right_margin = Cm(1.5)
    sec.different_first_page_header_footer = True  # титульник без номера
    footer = sec.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE \\* arabic \\* MERGEFORMAT")
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "24")
    rpr.append(sz)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = "2"
    r.append(t)
    fld.append(r)
    p._p.append(fld)


def title_par(doc, text, *, size=14, bold=False, before=0, after=0, align=WD_ALIGN_PARAGRAPH.CENTER):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = 1.15
    color = PLACEHOLDER_COLOR if is_placeholder(text) else BLACK
    add_run(p, text, bold=bold, size=size, color=color)
    return p


def build_title_page(doc, info, lab_no):
    title_par(doc, "Министерство образования Республики Беларусь")
    title_par(doc, "Учреждение образования")
    for line in ("«Белорусский государственный университет", "информатики и радиоэлектроники»"):
        title_par(doc, line)
    title_par(doc, info.get("department", "[Кафедра]"), before=28)
    title_par(doc, "ОТЧЁТ", size=16, bold=True, before=110)
    title_par(doc, f"по лабораторной работе № {lab_no}", before=6)
    title_par(doc, f"«{LAB_TITLES[lab_no]}»", bold=True)
    disc = info.get("discipline", "[Название дисциплины]")
    title_par(doc, f"по дисциплине «{disc}»")
    first_exec = True
    for label, value in (("Выполнил: студент гр. ", info.get("group", "[Номер группы]")),
                         ("", info.get("student", "[Фамилия И. О.]")),
                         ("Проверил: ", info.get("teacher", "[Фамилия И. О. преподавателя]"))):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf = p.paragraph_format
        pf.left_indent = Cm(8.5)
        pf.first_line_indent = Cm(0)
        pf.space_before = Pt(100 if first_exec else 2)
        first_exec = False
        if label:
            add_run(p, label, size=14)
        color = PLACEHOLDER_COLOR if is_placeholder(value) else BLACK
        add_run(p, value, size=14, color=color)
    title_par(doc, f"{info.get('city', 'Минск')} {info.get('year', '2026')}", before=140)
    doc.add_page_break()


def add_body_par(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(1.25)
    add_inline(p, text)
    return p


def add_code_par(doc, text, first=False, last=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.first_line_indent = Cm(0)
    pf.line_spacing = 1.0
    pf.space_before = Pt(6 if first else 0)
    pf.space_after = Pt(6 if last else 0)
    pf.left_indent = Cm(0.5)
    pf.right_indent = Cm(0.3)
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), "F2F2F2")
    pPr.append(shd)
    add_run(p, text if text else " ", mono=True, size=9, color=RGBColor(0x20, 0x20, 0x20))
    return p


def add_heading(doc, text, level):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Cm(0)
    add_inline(p, text, size=16 if level == 1 else 14, bold=True)
    return p


def add_table(doc, rows):
    ncols = max(len(r) for r in rows)
    tbl = doc.add_table(rows=len(rows), cols=ncols)
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = True
    for ri, row in enumerate(rows):
        for ci in range(ncols):
            cell = tbl.cell(ri, ci)
            text = row[ci] if ci < len(row) else ""
            par = cell.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT
            pf = par.paragraph_format
            pf.first_line_indent = Cm(0)
            pf.line_spacing = 1.0
            pf.space_before = Pt(2)
            pf.space_after = Pt(2)
            add_inline(par, text, size=11, bold=(ri == 0))
            tcPr = cell._tc.get_or_add_tcPr()
            mar = OxmlElement("w:tcMar")
            for side, w in (("top", 40), ("bottom", 40), ("start", 80), ("end", 80)):
                el = OxmlElement(f"w:{side}")
                el.set(qn("w:w"), str(w))
                el.set(qn("w:type"), "dxa")
                mar.append(el)
            tcPr.append(mar)
    # повтор шапки на новых страницах + запрет разрыва строк между страницами
    for row in tbl.rows:
        trPr = row._tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit")
        trPr.append(cs)
    trPr = tbl.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    trPr.append(th)
    return tbl


def split_table_row(line):
    # \| — экранированный пайп (литеральный символ), не разделитель колонок
    safe = "\x00"
    parts = line.strip().strip("|").replace("\\|", safe).split("|")
    return [c.strip().replace(safe, "|") for c in parts]


STRUCT_PREFIXES = ("```", "|", "#", "> ", "---", "***", "___", "- ", "* ")


def is_structural(s):
    return (not s) or s.startswith(STRUCT_PREFIXES) or re.match(r"^\d{1,2}[.)]\s", s)


def convert_md(md_text, doc):
    lines = md_text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if stripped.startswith("```"):
            i += 1
            buf = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                buf.append(lines[i].rstrip())
                i += 1
            i += 1  # закрывающая ```
            for k, code_line in enumerate(buf[:400]):
                add_code_par(doc, code_line, first=(k == 0), last=(k == len(buf[:400]) - 1))
            if len(buf) > 400:
                add_code_par(doc, f"... (усечено, всего строк: {len(buf)})", last=True)
            continue
        if stripped.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:\-|]+\|?$", lines[i + 1].strip()):
            rows = [split_table_row(stripped)]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(split_table_row(lines[i].strip()))
                i += 1
            add_table(doc, rows)
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
            continue
        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            add_heading(doc, m.group(2).strip(), min(len(m.group(1)), 3))
            i += 1
            continue
        if stripped in ("---", "***", "___"):
            i += 1
            continue
        if not stripped:
            i += 1
            continue
        # абзац: склеиваем жёстко перенесённые строки markdown до пустой/структурной строки
        chunk = [stripped]
        i += 1
        while i < len(lines) and not is_structural(lines[i].strip()):
            chunk.append(lines[i].strip())
            i += 1
        text = " ".join(chunk)
        if text.startswith("> "):
            p = add_body_par(doc, text[2:])
            for r in p.runs:
                r.italic = True
                r.font.color.rgb = GRAY
        elif stripped.startswith(("- ", "* ")):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.left_indent = Cm(1.0)
            add_inline(p, "–\u00A0" + text[2:])
        else:
            m = re.match(r"^(\d{1,2})[.)]\s+(.*)$", text)
            if m:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                p.paragraph_format.first_line_indent = Cm(0)
                p.paragraph_format.left_indent = Cm(1.0)
                add_inline(p, f"{m.group(1)}.\u00A0{m.group(2)}")
            else:
                add_body_par(doc, text)


def build_lab(lab_no, info):
    md_path = BASE / f"lab_{lab_no}" / "report" / "report.md"
    if not md_path.exists():
        print(f"lab_{lab_no}: report.md не найден, пропуск")
        return False
    doc = Document()
    setup_styles(doc)
    setup_page(doc)
    build_title_page(doc, info, lab_no)
    convert_md(md_path.read_text(encoding="utf-8"), doc)
    out = BASE / f"lab_{lab_no}" / "report" / f"Отчет_ЛР_{lab_no}.docx"
    doc.save(out)
    print(f"lab_{lab_no}: OK -> {out}")
    return True


def main():
    info = json.loads((BUILD / "student_info.json").read_text(encoding="utf-8"))
    args = [a for a in sys.argv[1:] if a.startswith("lab_")]
    if args:
        nums = [int(a.split("_")[1]) for a in args]
    else:
        nums = sorted(n for n in LAB_TITLES
                      if (BASE / f"lab_{n}" / "report" / "report.md").exists())
    built = [n for n in nums if build_lab(n, info)]
    print(f"Собрано отчётов: {len(built)} из {len(nums)}: {built}")


if __name__ == "__main__":
    main()
